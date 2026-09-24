#!/usr/bin/env python3
"""SQLite-Datenschicht für die Jägerprüfungs-Fragendatenbank.

Enthält:
  - init_db(conn): legt Tabellen an (falls nicht vorhanden)
  - import_excel(conn, path): liest Jagdfragen_*.xlsx (Sheet 'Fragen') ein
    und aktualisiert die Datenbank (Upsert anhand des Fragetexts)
  - Abfragefunktionen für die GUI (Fragen, Kategorien, Statistik)
"""
import os
import shutil
import sqlite3
from pathlib import Path

import openpyxl

DB_FILENAME = "jagdfragen.db"

SCHEMA = """
CREATE TABLE IF NOT EXISTS questions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    frage TEXT NOT NULL UNIQUE,
    erlaeuterung TEXT DEFAULT '',
    kategorie TEXT DEFAULT '',
    bild TEXT DEFAULT '',
    typ TEXT DEFAULT 'mc',
    audio TEXT DEFAULT ''
);

CREATE TABLE IF NOT EXISTS answers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    question_id INTEGER NOT NULL REFERENCES questions(id) ON DELETE CASCADE,
    letter TEXT NOT NULL CHECK(letter IN ('A','B','C','D')),
    text TEXT NOT NULL DEFAULT '',
    is_correct INTEGER NOT NULL DEFAULT 0
);

CREATE TABLE IF NOT EXISTS stats (
    question_id INTEGER PRIMARY KEY REFERENCES questions(id) ON DELETE CASCADE,
    mal_gezeigt INTEGER NOT NULL DEFAULT 0,
    mal_richtig INTEGER NOT NULL DEFAULT 0,
    zuletzt_falsch INTEGER NOT NULL DEFAULT 0
);
"""


def get_connection(db_path=DB_FILENAME):
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db(conn):
    conn.executescript(SCHEMA)
    conn.commit()
    _migrate_add_column(conn, "bild", "TEXT DEFAULT ''")
    _migrate_add_column(conn, "typ", "TEXT DEFAULT 'mc'")
    _migrate_add_column(conn, "audio", "TEXT DEFAULT ''")


def _migrate_add_column(conn, spalte, definition):
    """Fügt eine Spalte nachträglich zur Tabelle 'questions' hinzu, falls die
    Datenbank aus einer älteren App-Version stammt und sie noch fehlt."""
    spalten = [r["name"] for r in conn.execute("PRAGMA table_info(questions)").fetchall()]
    if spalte not in spalten:
        conn.execute(f"ALTER TABLE questions ADD COLUMN {spalte} {definition}")
        conn.commit()


class ImportError_(Exception):
    pass


def _read_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Fragen" not in wb.sheetnames:
        raise ImportError_(
            f"Die Datei enthält kein Tabellenblatt 'Fragen' (gefunden: {wb.sheetnames})."
        )
    ws = wb["Fragen"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return rows


def excel_fragen_set(path):
    """Liest nur die Fragetexte (Spalte 'Frage') aus einer Excel-Datei ein,
    ohne die Datenbank zu verändern. Nützlich, um vor einem Sync zu prüfen,
    welche vorhandenen Fragen in der Datei fehlen (und damit beim
    Synchronisieren entfernt würden)."""
    path = Path(path)
    if not path.exists():
        raise ImportError_(f"Datei nicht gefunden: {path}")
    rows = _read_rows(path)
    frages = set()
    for row in rows:
        if row is None or all(v in (None, "") for v in row):
            continue
        frage = (row[0] or "").strip()
        if frage:
            frages.add(frage)
    return frages


def preview_sync_excel(conn, path):
    """Vergleicht die Fragen in der Excel-Datei mit der Datenbank, OHNE etwas
    zu ändern. Rückgabe: dict mit 'werden_entfernt' (Liste von Fragetexten,
    die in der DB stehen, aber nicht mehr in der Excel-Datei vorkommen) und
    'anzahl_excel' (Gesamtzahl Fragen in der Excel-Datei)."""
    excel_frages = excel_fragen_set(path)
    db_frages = {r["frage"] for r in conn.execute("SELECT frage FROM questions").fetchall()}
    werden_entfernt = sorted(db_frages - excel_frages)
    return {"werden_entfernt": werden_entfernt, "anzahl_excel": len(excel_frages)}


def import_excel(conn, path, bilder_zielordner=None, audio_zielordner=None, delete_missing=False):
    """Importiert/aktualisiert Fragen aus einer Excel-Datei.

    Falls eine Zeile in der Spalte 'Bild' einen Dateinamen enthält, wird die
    zugehörige Bilddatei aus dem Unterordner 'bilder' neben der Excel-Datei
    gesucht und (wenn bilder_zielordner angegeben ist) dorthin kopiert, damit
    die App sie unabhängig vom Speicherort der Excel-Datei findet. Analog
    dazu wird eine Datei in der Spalte 'Audio' aus dem Unterordner 'audio'
    gesucht und (wenn audio_zielordner angegeben ist) dorthin kopiert.

    Ist delete_missing=True (Synchronisieren statt reinem Import), werden am
    Ende zusätzlich alle Fragen aus der Datenbank gelöscht, deren Fragetext
    nicht (mehr) in der Excel-Datei vorkommt – zusammen mit ihren Antworten
    und ihrer Statistik (ON DELETE CASCADE). Standardmäßig (False) wird wie
    bisher nie gelöscht, nur ergänzt/aktualisiert.

    Rückgabe: dict mit Zählern (neu, aktualisiert, uebersprungen, geloescht-
    Liste, fehler-Liste).
    """
    path = Path(path)
    if not path.exists():
        raise ImportError_(f"Datei nicht gefunden: {path}")

    rows = _read_rows(path)
    bilder_quellordner = path.parent / "bilder"
    audio_quellordner = path.parent / "audio"

    neu = 0
    aktualisiert = 0
    uebersprungen = 0
    fehler = []
    gesehene_fragen = set()

    cur = conn.cursor()
    for idx, row in enumerate(rows, start=2):  # Excel-Zeilennummer für Fehlermeldungen
        if row is None or all(v in (None, "") for v in row):
            continue

        # Auf 14 Spalten auffüllen, falls Zeile kürzer ist (Audio ist die
        # neue, optionale 14. Spalte – ans Ende gehängt, damit bestehende
        # Excel-Dateien nicht umsortiert werden müssen)
        row = list(row) + [None] * (14 - len(row))
        (frage, bild, a_a, a_b, a_c, a_d, r_a, r_b, r_c, r_d, erlaeuterung, kategorie, fragetyp, audio) = row[:14]

        frage = (frage or "").strip()
        if not frage:
            uebersprungen += 1
            continue

        typ = "text" if (fragetyp or "").strip().lower().startswith("text") else "mc"

        bild = (bild or "").strip()
        if bild:
            quelle = bilder_quellordner / bild
            if not quelle.exists():
                fehler.append(
                    f"Zeile {idx}: Bilddatei '{bild}' nicht gefunden in '{bilder_quellordner}' "
                    "– Frage wird ohne Bild importiert."
                )
                bild = ""
            elif bilder_zielordner:
                os.makedirs(bilder_zielordner, exist_ok=True)
                ziel = Path(bilder_zielordner) / bild
                if quelle.resolve() != ziel.resolve():
                    shutil.copyfile(quelle, ziel)

        audio = (audio or "").strip()
        if audio:
            quelle_audio = audio_quellordner / audio
            if not quelle_audio.exists():
                fehler.append(
                    f"Zeile {idx}: Audiodatei '{audio}' nicht gefunden in '{audio_quellordner}' "
                    "– Frage wird ohne Audio importiert."
                )
                audio = ""
            elif audio_zielordner:
                os.makedirs(audio_zielordner, exist_ok=True)
                ziel_audio = Path(audio_zielordner) / audio
                if quelle_audio.resolve() != ziel_audio.resolve():
                    shutil.copyfile(quelle_audio, ziel_audio)

        antworten = {
            "A": (a_a or "").strip(),
            "B": (a_b or "").strip(),
            "C": (a_c or "").strip(),
            "D": (a_d or "").strip(),
        }
        richtig_raw = {"A": r_a, "B": r_b, "C": r_c, "D": r_d}

        def ist_richtig(v):
            if v is None:
                return False
            s = str(v).strip().lower()
            return s in ("x", "wahr", "true", "1", "ja")

        richtig = {k: ist_richtig(v) for k, v in richtig_raw.items()}

        gefuellte_antworten = [k for k, v in antworten.items() if v]
        if typ == "text":
            if not any(richtig.get(k) for k in gefuellte_antworten):
                fehler.append(
                    f"Zeile {idx}: Text-Frage ohne als richtig markierte Antwort (Antwort_A + Richtig_A "
                    "ausfüllen) – übersprungen."
                )
                uebersprungen += 1
                continue
        else:
            if len(gefuellte_antworten) < 2:
                fehler.append(f"Zeile {idx}: weniger als 2 Antwortmöglichkeiten ausgefüllt – übersprungen.")
                uebersprungen += 1
                continue
            if not any(richtig.get(k) for k in gefuellte_antworten):
                fehler.append(f"Zeile {idx}: keine Antwort als richtig markiert – übersprungen.")
                uebersprungen += 1
                continue

        erlaeuterung = (erlaeuterung or "").strip()
        kategorie = (kategorie or "").strip()
        gesehene_fragen.add(frage)

        existing = cur.execute(
            "SELECT id FROM questions WHERE frage = ?", (frage,)
        ).fetchone()

        if existing:
            qid = existing["id"]
            cur.execute(
                "UPDATE questions SET erlaeuterung = ?, kategorie = ?, bild = ?, typ = ?, audio = ? WHERE id = ?",
                (erlaeuterung, kategorie, bild, typ, audio, qid),
            )
            cur.execute("DELETE FROM answers WHERE question_id = ?", (qid,))
            aktualisiert += 1
        else:
            cur.execute(
                "INSERT INTO questions (frage, erlaeuterung, kategorie, bild, typ, audio) VALUES (?, ?, ?, ?, ?, ?)",
                (frage, erlaeuterung, kategorie, bild, typ, audio),
            )
            qid = cur.lastrowid
            cur.execute(
                "INSERT OR IGNORE INTO stats (question_id) VALUES (?)", (qid,)
            )
            neu += 1

        for letter in ("A", "B", "C", "D"):
            text = antworten[letter]
            if not text:
                continue
            cur.execute(
                "INSERT INTO answers (question_id, letter, text, is_correct) VALUES (?, ?, ?, ?)",
                (qid, letter, text, 1 if richtig[letter] else 0),
            )

    geloescht = []
    if delete_missing:
        vorhandene = cur.execute("SELECT id, frage FROM questions").fetchall()
        for row in vorhandene:
            if row["frage"] not in gesehene_fragen:
                cur.execute("DELETE FROM questions WHERE id = ?", (row["id"],))
                geloescht.append(row["frage"])

    conn.commit()
    return {
        "neu": neu,
        "aktualisiert": aktualisiert,
        "uebersprungen": uebersprungen,
        "geloescht": geloescht,
        "fehler": fehler,
        "gesamt": neu + aktualisiert,
    }


def list_categories(conn):
    rows = conn.execute(
        "SELECT DISTINCT kategorie FROM questions WHERE kategorie != '' ORDER BY kategorie"
    ).fetchall()
    return [r["kategorie"] for r in rows]


def get_questions(conn, kategorie=None, nur_falsche=False):
    """Liefert Fragen (mit Antworten) als Liste von dicts.

    kategorie kann eine einzelne Kategorie (String), eine Liste/ein Set
    mehrerer Kategorien, oder None/leer sein (= keine Einschränkung, alle
    Kategorien)."""
    sql = "SELECT * FROM questions"
    params = []
    clauses = []
    if kategorie:
        if isinstance(kategorie, (list, tuple, set)):
            gewaehlt = [k for k in kategorie if k]
            if gewaehlt:
                platzhalter = ",".join("?" for _ in gewaehlt)
                clauses.append(f"kategorie IN ({platzhalter})")
                params.extend(gewaehlt)
        else:
            clauses.append("kategorie = ?")
            params.append(kategorie)
    if nur_falsche:
        clauses.append(
            "id IN (SELECT question_id FROM stats WHERE zuletzt_falsch = 1)"
        )
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY id"

    frage_rows = conn.execute(sql, params).fetchall()
    result = []
    for fr in frage_rows:
        antworten = conn.execute(
            "SELECT letter, text, is_correct FROM answers WHERE question_id = ? ORDER BY letter",
            (fr["id"],),
        ).fetchall()
        result.append(
            {
                "id": fr["id"],
                "frage": fr["frage"],
                "erlaeuterung": fr["erlaeuterung"],
                "kategorie": fr["kategorie"],
                "bild": fr["bild"] or "",
                "typ": fr["typ"] or "mc",
                "audio": fr["audio"] or "",
                "antworten": [
                    {"letter": a["letter"], "text": a["text"], "is_correct": bool(a["is_correct"])}
                    for a in antworten
                ],
            }
        )
    return result


def record_attempt(conn, question_id, correct):
    conn.execute(
        "INSERT INTO stats (question_id, mal_gezeigt, mal_richtig, zuletzt_falsch) "
        "VALUES (?, 1, ?, ?) "
        "ON CONFLICT(question_id) DO UPDATE SET "
        "mal_gezeigt = mal_gezeigt + 1, "
        "mal_richtig = mal_richtig + excluded.mal_richtig, "
        "zuletzt_falsch = excluded.zuletzt_falsch",
        (question_id, 1 if correct else 0, 0 if correct else 1),
    )
    conn.commit()


def question_count(conn):
    return conn.execute("SELECT COUNT(*) AS c FROM questions").fetchone()["c"]
